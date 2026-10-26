"""Tests for TidyTable; all samples are generated in tmp_path."""
import csv

import pytest
from openpyxl import Workbook, load_workbook

from tidytable.cli import main
from tidytable.core import parse_number


def read_csv(p):
    with p.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def make(p, text):
    p.write_text(text, encoding="utf-8")
    return p


def test_trim_case_and_exact_dupes(tmp_path, capsys):
    src = make(tmp_path / "in.csv", 'name,city\n"  alice   SMITH "," paris"\nAlice Smith,Paris\nBOB jones,london\n')
    out = tmp_path / "out.csv"
    assert main([str(src), "-o", str(out), "--case", "title:name,city"]) == 0
    assert read_csv(out) == [{"name": "Alice Smith", "city": "Paris"}, {"name": "Bob Jones", "city": "London"}]
    assert "duplicates removed: 1" in capsys.readouterr().out


def test_numeric_repair_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["item", "price"])
    for r in [("a", "$1,234.50"), ("b", "1.234,50 €"), ("c", "12"), ("d", "USD 7")]:
        ws.append(r)
    src = tmp_path / "in.xlsx"
    wb.save(src)
    out = tmp_path / "out.xlsx"
    assert main([str(src), "-o", str(out)]) == 0
    vals = [r[1] for r in load_workbook(out).active.iter_rows(min_row=2, values_only=True)]
    assert vals == [1234.5, 1234.5, 12, 7]


def test_required_and_no_negative(tmp_path):
    src = make(tmp_path / "in.csv", "product,qty\nWidget,5\nGadget,-3\n,4\nBolt,\n")
    out, rep = tmp_path / "out.csv", tmp_path / "rep.csv"
    rc = main([str(src), "-o", str(out), "--required", "product,qty", "--no-negative", "qty", "--report", str(rep)])
    assert rc == 0
    assert read_csv(out) == [{"product": "Widget", "qty": "5"}]
    by_row = {r["row"]: r for r in read_csv(rep)}
    assert set(by_row) == {"3", "4", "5"}
    assert "negative value in 'qty'" in by_row["3"]["reason"]
    assert "missing required value in 'product'" in by_row["4"]["reason"]
    assert "missing required value in 'qty'" in by_row["5"]["reason"]
    assert all(r["action"] == "removed" for r in by_row.values())


def test_ambiguous_number_flagged_not_guessed(tmp_path):
    src = make(tmp_path / "in.csv", "item,price\na,10\nb,abc\nc,$5\n")
    out, rep = tmp_path / "out.csv", tmp_path / "rep.csv"
    assert main([str(src), "-o", str(out), "--report", str(rep)]) == 0
    assert [r["price"] for r in read_csv(out)] == ["10", "abc", "5"]
    flagged = read_csv(rep)
    assert len(flagged) == 1 and flagged[0]["action"] == "kept for review"
    assert "abc" in flagged[0]["reason"]


def test_key_dedupe_ignores_case_and_spaces(tmp_path):
    src = make(tmp_path / "in.csv", 'email,name\n" Bob@X.com ",Bob\nbob@x.com,Robert\ncarol@y.com,Carol\n')
    out = tmp_path / "out.csv"
    assert main([str(src), "-o", str(out), "--key", "email"]) == 0
    assert read_csv(out) == [{"email": "Bob@X.com", "name": "Bob"}, {"email": "carol@y.com", "name": "Carol"}]


def test_bad_input_errors(tmp_path, capsys):
    assert main([str(tmp_path / "nope.csv"), "-o", str(tmp_path / "o.csv")]) != 0
    assert "not found" in capsys.readouterr().err
    src = make(tmp_path / "in.csv", "a,b\n1,2\n")
    assert main([str(src), "-o", str(tmp_path / "o.csv"), "--key", "zzz"]) != 0
    assert "unknown column" in capsys.readouterr().err
    assert not (tmp_path / "o.csv").exists()


def test_input_never_modified(tmp_path, capsys):
    src = make(tmp_path / "in.csv", "a,b\n x ,2\n")
    before = src.read_bytes()
    assert main([str(src), "-o", str(src)]) != 0
    assert "overwrite" in capsys.readouterr().err
    assert src.read_bytes() == before
    assert main([str(src), "-o", str(tmp_path / "out.csv")]) == 0
    assert src.read_bytes() == before


def test_csv_to_xlsx_roundtrip(tmp_path):
    src = make(tmp_path / "in.csv", "a,b\n1,x\n2,y\n3,z\n")
    out = tmp_path / "out.xlsx"
    assert main([str(src), "-o", str(out)]) == 0
    assert load_workbook(out).active.max_row == 4


@pytest.mark.parametrize("text,expected", [
    ("1.234,56", 1234.56), ("1,234.56", 1234.56), ("$ 12", 12), ("(5)", -5),
    ("1,234", 1234), ("0.125", 0.125), ("abc", None), ("1.234", None), ("12kg", None),
])
def test_parse_number(text, expected):
    assert parse_number(text) == pytest.approx(expected) if expected is not None else parse_number(text) is None


def test_lenient_letters():
    assert parse_number("12kg", lenient=True) == 12
    assert parse_number("a1b2", lenient=True) is None
